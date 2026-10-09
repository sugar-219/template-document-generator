"""Local data import and loss-minimizing DOCX mail merge."""
from __future__ import annotations

import csv
import copy
import hashlib
import io
import re
import shutil
import threading
import uuid
import zipfile
from bisect import bisect_right
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Callable

from lxml import etree
from openpyxl import load_workbook
from pdf_worker import ConversionCancelled, PdfSession

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
TAG = lambda name: f"{{{W}}}{name}"
TOKEN = re.compile(r"\{\{\s*([^{}\r\n\t\x00]+?)\s*\}\}")
PART = re.compile(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


class InputError(ValueError):
    """A problem the user can correct in a source file or setting."""


@dataclass
class Dataset:
    headers: list[str]
    records: list[dict[str, str]]
    row_numbers: list[int]
    sheet: str


def cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "e":
        raise InputError(f"单元格 {cell.coordinate} 为 Excel 错误值 {value}，请先修正。")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d" if not value.time().isoformat().strip("0:") else "%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (float, int)):
        fmt = cell.number_format.split(";")[0]
        if re.fullmatch(r"0+", fmt) and float(value).is_integer():
            return str(int(value)).zfill(len(fmt))
        if re.fullmatch(r"0\.0+", fmt):
            return f"{value:.{len(fmt.split('.')[1])}f}"
        if re.fullmatch(r"0(?:\.0+)?%", fmt):
            decimals = len(fmt.split('.')[1].rstrip('%')) if "." in fmt else 0
            return f"{value * 100:.{decimals}f}%"
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def _dataset(rows, sheet: str) -> Dataset:
    rows = iter(rows)
    first = next(rows, None)
    if first is None:
        raise InputError("数据文件为空，请在第一行填写列标题。")
    raw_headers = [v.strip() for v in first]
    active = [i for i, value in enumerate(raw_headers) if value]
    if not active:
        raise InputError("第一行没有列标题。请将标题放在第一行。")
    headers = [raw_headers[i] for i in active]
    duplicate = sorted({h for h in headers if headers.count(h) > 1})
    if duplicate:
        raise InputError("列标题重复：" + "、".join(duplicate))
    records, numbers = [], []
    for row_number, row in enumerate(rows, 2):
        if not any(v.strip() for v in row):
            continue
        for index, value in enumerate(row):
            if value.strip() and (index >= len(raw_headers) or not raw_headers[index]):
                raise InputError(f"第 {row_number} 行第 {index + 1} 列有数据但没有列标题。")
        records.append({h: row[i] if i < len(row) else "" for h, i in zip(headers, active)})
        numbers.append(row_number)
    if not records:
        raise InputError("只有标题，没有可生成的数据记录。")
    return Dataset(headers, records, numbers, sheet)


def sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() == ".csv":
        return ["CSV"]
    if path.suffix.lower() != ".xlsx":
        raise InputError("仅支持 .xlsx 或 .csv，请将旧版 .xls 另存为 .xlsx。")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_data(path: Path, sheet: str | None = None) -> Dataset:
    if path.suffix.lower() == ".csv":
        content = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise InputError("CSV 编码无法识别，请另存为 UTF-8 CSV。")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return _dataset(csv.reader(io.StringIO(text, newline=""), dialect), "CSV")
    if path.suffix.lower() != ".xlsx":
        raise InputError("请选择 .xlsx 或 .csv 数据文件。")
    book = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    try:
        name = sheet or book.sheetnames[0]
        if name not in book.sheetnames:
            raise InputError(f"找不到工作表：{name}")
        def rows():
            for formula_row, value_row in zip(book[name].iter_rows(), cached[name].iter_rows()):
                row = []
                for source, value in zip(formula_row, value_row):
                    if source.data_type == "f" and value.value is None:
                        raise InputError(f"单元格 {source.coordinate} 的公式没有缓存结果。请用 Excel 或 WPS 打开、计算并保存后重新导入。")
                    row.append(cell_text(value))
                yield row
        return _dataset(rows(), name)
    finally:
        book.close()
        cached.close()


def paragraph_nodes(paragraph):
    """Keep nested text-box paragraphs separate, with barriers at tabs/line breaks."""
    nodes, positions, chunks, offset = [], [], [], 0
    for node in paragraph.iter():
        if node.tag not in (TAG("t"), TAG("br"), TAG("tab"), TAG("cr"), TAG("fldChar")):
            continue
        if next(node.iterancestors(TAG("p")), None) is not paragraph:
            continue
        if any(ancestor.tag == TAG("del") for ancestor in node.iterancestors()):
            continue
        if node.tag == TAG("t"):
            text = node.text or ""
            if text:
                nodes.append(node)
                positions.append(offset)
                chunks.append(text)
                offset += len(text)
        else:
            chunks.append("\x00")
            offset += 1
    return nodes, positions, "".join(chunks)


def _write_text(node, value: str):
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    pieces = re.split(r"([\n\t])", value)
    node.text = pieces[0]
    node.set(XML_SPACE, "preserve")
    parent = node.getparent()
    position = parent.index(node) + 1
    for piece in pieces[1:]:
        if piece in ("\n", "\t"):
            item = etree.Element(TAG("br" if piece == "\n" else "tab"))
        else:
            item = etree.Element(TAG("t"))
            item.text = piece
            item.set(XML_SPACE, "preserve")
        parent.insert(position, item)
        position += 1


def replace_root(root, values: dict[str, str]):
    for paragraph in root.iter(TAG("p")):
        nodes, positions, text = paragraph_nodes(paragraph)
        changed = set()
        for match in reversed(list(TOKEN.finditer(text))):
            field = match.group(1).strip()
            if field not in values:
                continue
            first = bisect_right(positions, match.start()) - 1
            last = bisect_right(positions, match.end() - 1) - 1
            start, end = match.start() - positions[first], match.end() - positions[last]
            if first == last:
                old = nodes[first].text or ""
                nodes[first].text = old[:start] + values[field] + old[end:]
            else:
                nodes[first].text = (nodes[first].text or "")[:start] + values[field]
                nodes[last].text = (nodes[last].text or "")[end:]
                for index in range(first + 1, last):
                    nodes[index].text = ""
                changed.update(range(first + 1, last + 1))
            changed.add(first)
        for index in sorted(changed):
            _write_text(nodes[index], nodes[index].text or "")


class Template:
    def __init__(self, path: Path):
        self.path = Path(path)
        if self.path.suffix.lower() != ".docx":
            raise InputError("模板必须是 .docx，请将 .doc 或 .docm 另存为 .docx。")
        try:
            with zipfile.ZipFile(self.path) as archive:
                self.entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
        except zipfile.BadZipFile as exc:
            raise InputError("模板不是有效的 DOCX，或文件已加密。请打开后另存为未加密 .docx。") from exc
        self.parts = {}
        self.roots = {}
        self.fingerprint = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.fields = []
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        for item, data in self.entries:
            if PART.fullmatch(item.filename):
                root = etree.fromstring(data, parser)
                self.parts[item.filename] = data
                self.roots[item.filename] = root
                for paragraph in root.iter(TAG("p")):
                    text = paragraph_nodes(paragraph)[2]
                    for match in TOKEN.finditer(text):
                        field = match.group(1).strip()
                        if field and field not in self.fields:
                            self.fields.append(field)
        if "word/document.xml" not in self.parts:
            raise InputError("DOCX 缺少正文，无法使用。")
        if not self.fields:
            raise InputError("模板中未找到 {{字段名}}。请在正文、表格或页眉页脚中添加占位符并保存。")

    def filled_roots(self, values: dict[str, str]):
        missing = [field for field in self.fields if field not in values]
        if missing:
            raise InputError("这些模板字段尚未匹配：" + "、".join(missing))
        if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value) for value in values.values()):
            raise InputError("数据包含 Word 不支持的控制字符，请清理后重试。")
        output = {}
        for name, original in self.roots.items():
            root = copy.deepcopy(original)
            replace_root(root, values)
            output[name] = root
        return output

    def rendered_parts(self, values: dict[str, str]) -> dict[str, bytes]:
        return {name: etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                for name, root in self.filled_roots(values).items()}

    def render(self, values: dict[str, str], destination: Path):
        parts = self.rendered_parts(values)
        destination = Path(destination)
        # Exclusive create means the source or an existing result can never be overwritten.
        with destination.open("xb") as stream:
            try:
                with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
                    for item, data in self.entries:
                        archive.writestr(item, parts.get(item.filename, data))
            except Exception:
                stream.close()
                destination.unlink(missing_ok=True)
                raise

    def preview(self, values: dict[str, str]) -> str:
        blocks = []
        for name, root in self.filled_roots(values).items():
            label = "正文" if name == "word/document.xml" else ("页眉" if "header" in name else "页脚" if "footer" in name else "脚注 / 尾注")
            paragraphs = []
            for paragraph in root.iter(TAG("p")):
                text = paragraph_nodes(paragraph)[2].replace("\x00", "\n")
                if text.strip():
                    paragraphs.append(text)
            if paragraphs:
                blocks.append(f"【{label}】\n\n" + "\n".join(paragraphs))
        return "\n\n".join(blocks)


def values_for(template: Template, mapping: dict[str, str], record: dict[str, str]) -> dict[str, str]:
    missing = [field for field in template.fields if mapping.get(field) not in record]
    if missing:
        raise InputError("请选择这些字段对应的数据列：" + "、".join(missing))
    return {field: record[mapping[field]] for field in template.fields}


def file_stem(pattern: str, values: dict[str, str], sequence: int) -> str:
    if not pattern.strip():
        raise InputError("请填写文件命名规则。")
    unknown = [m.group(1).strip() for m in TOKEN.finditer(pattern) if m.group(1).strip() not in values and m.group(1).strip() != "序号"]
    if unknown:
        raise InputError("命名规则中的字段不存在：" + "、".join(unknown))
    all_values = {**values, "序号": f"{sequence:03d}"}
    text = TOKEN.sub(lambda match: all_values[match.group(1).strip()], pattern)
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text).strip().rstrip(" .")[:100].rstrip(" .")
    if not text:
        text = f"文档_{sequence:03d}"
    if text.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        text = "_" + text
    return text


def preview_filenames(template, dataset, mapping, pattern, indices=None, limit=5):
    indices = range(len(dataset.records)) if indices is None else indices
    items, used = [], set()
    for sequence, index in enumerate(indices[:limit], 1):
        values = values_for(template, mapping, dataset.records[index])
        original = file_stem(pattern, values, sequence)
        stem, suffix = original, 2
        while stem.casefold() in used:
            stem, suffix = f"{original}_{suffix}", suffix + 1
        used.add(stem.casefold())
        fields = dict.fromkeys(match.group(1).strip() for match in TOKEN.finditer(pattern))
        empty = [name for name in fields if name != "序号" and not values[name].strip()]
        raw = TOKEN.sub(lambda match: {**values, "序号": f"{sequence:03d}"}[match.group(1).strip()], pattern)
        notes = (["空值字段：" + "、".join(empty)] if empty else [])
        if raw != original:
            notes.append("名称已按 Windows 文件名规则调整")
        if stem != original:
            notes.append("重名，已追加编号")
        items.append({"数据行": str(dataset.row_numbers[index]), "Word 文件": stem + ".docx", "说明": "；".join(notes)})
    return items


def pdf_converter() -> tuple[str, str] | None:
    import os
    if os.name == "nt":
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Word.Application\CLSID"):
                return ("Word", "")
        except OSError:
            pass
    candidates = [shutil.which("soffice")]
    for key in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
        base = os.environ.get(key)
        if base:
            candidates.append(str(Path(base) / "LibreOffice/program/soffice.exe"))
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return ("LibreOffice", candidate)
    return None


@dataclass
class BatchResult:
    folder: Path
    items: list[dict[str, str]]
    cancelled: bool
    indices: list[int] = field(default_factory=list)
    export_pdf: bool = False
    manifest_path: Path | None = None
    manifest_warning: str = ""
    remaining_indices: list[int] = field(default_factory=list)
    used_stems: set[str] = field(default_factory=set)

    @property
    def succeeded(self):
        return sum(item["状态"] == "成功" for item in self.items)

    @property
    def partial(self):
        return sum(item["状态"] == "部分成功" for item in self.items)


def check_records(template, dataset, mapping, pattern, indices=None):
    indices = range(len(dataset.records)) if indices is None else indices
    issues, seen = [], set()
    if any(mapping.get(name) not in dataset.headers for name in template.fields):
        raise InputError("请先完成所有字段匹配。")
    for sequence, index in enumerate(indices, 1):
        record = dataset.records[index]
        values = values_for(template, mapping, record)
        for name, value in values.items():
            reason = "包含 Word 不支持的控制字符" if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value) else "空值将按空白填充" if not value.strip() else ""
            if reason:
                issues.append({"index": index, "行号": dataset.row_numbers[index], "字段": name, "数据列": mapping[name],
                               "级别": "错误" if "控制字符" in reason else "提示", "说明": reason})
        stem = file_stem(pattern, values, sequence).casefold()
        if stem in seen:
            issues.append({"index": index, "行号": dataset.row_numbers[index], "字段": "文件命名", "数据列": "", "级别": "提示", "说明": "文件名重复，将自动添加编号"})
        seen.add(stem)
    return issues


def save_manifest(result):
    target = result.folder / "生成清单.csv"
    temporary = target.with_suffix(".tmp")
    result.manifest_path, result.manifest_warning = None, ""
    try:
        with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=["数据行", "Word 文件", "PDF 文件", "状态", "说明"])
            writer.writeheader()
            writer.writerows(result.items)
        try:
            temporary.replace(target)
        except OSError:
            target = result.folder / (datetime.now().strftime("生成清单_%Y%m%d_%H%M%S_") + uuid.uuid4().hex + ".csv")
            temporary.rename(target)
            result.manifest_warning = "生成清单.csv 无法更新，已另存为 " + target.name + "。"
        result.manifest_path = target
    except OSError as exc:
        result.manifest_warning = "生成清单未能保存：" + str(exc) + "。文档处理结果已保留，请检查文件占用或输出目录后重试。"
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def generate_batch(template: Template, dataset: Dataset, mapping: dict[str, str], output: Path,
                   pattern: str, indices: list[int] | None = None, export_pdf: bool = False,
                   cancel: threading.Event | None = None,
                   progress: Callable[[int, int, dict[str, str]], None] | None = None,
                   *, previous: BatchResult | None = None,
                   phase: Callable[[str, int, int], None] | None = None) -> BatchResult:
    indices = list(range(len(dataset.records))) if indices is None else list(indices)
    if not indices:
        raise InputError("没有选择要生成的记录。")
    if any(index < 0 or index >= len(dataset.records) for index in indices):
        raise InputError("选择的记录已失效，请重新选择。")
    first_values = values_for(template, mapping, dataset.records[indices[0]])
    initial_count = len(previous.items) if previous else 0
    file_stem(pattern, first_values, initial_count + 1)
    converter = pdf_converter() if export_pdf else None
    if previous:
        folder = previous.folder
        result = BatchResult(folder, [dict(item) for item in previous.items], False,
                             list(previous.indices), export_pdf, used_stems=set(previous.used_stems))
        result.used_stems.update(path.stem.casefold() for path in folder.iterdir()
                                 if path.suffix.lower() in (".docx", ".pdf"))
    else:
        output = Path(output)
        output.mkdir(parents=True, exist_ok=True)
        folder = output / (datetime.now().strftime("批次_%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:6])
        folder.mkdir()
        result = BatchResult(folder, [], False, [], export_pdf)
    result.remaining_indices = list(indices)
    session = None
    try:
        if converter:
            if phase:
                phase("准备转换器", 0, len(indices))
            session = PdfSession(converter, cancel=cancel).__enter__()
        used = result.used_stems
        for number, index in enumerate(indices, 1):
            if cancel and cancel.is_set():
                result.cancelled = True
                break
            item = {"数据行": str(dataset.row_numbers[index]), "Word 文件": "", "PDF 文件": "", "状态": "失败", "说明": ""}
            try:
                values = values_for(template, mapping, dataset.records[index])
                original = file_stem(pattern, values, initial_count + number)
                stem, suffix = original, 2
                while stem.casefold() in used:
                    stem, suffix = f"{original}_{suffix}", suffix + 1
                used.add(stem.casefold())
                destination = folder / f"{stem}.docx"
                if phase:
                    phase("填充 Word", number, len(indices))
                template.render(values, destination)
                item["Word 文件"] = destination.name
                if export_pdf and not session:
                    raise InputError("未检测到可用 PDF 转换器。")
                if session:
                    if phase:
                        phase("转换 PDF", number, len(indices))
                    item["PDF 文件"] = session.convert(destination).name
                item["状态"] = "成功"
                item["说明"] = "含空值字段" if any(not value.strip() for value in values.values()) else ""
            except Exception as exc:
                item["说明"] = str(exc)
                if item["Word 文件"] and export_pdf:
                    item["状态"] = "部分成功"
                    item["说明"] = "Word 已生成；PDF 失败：" + str(exc)
                if isinstance(exc, ConversionCancelled):
                    result.cancelled = True
            result.items.append(item)
            result.indices.append(index)
            if progress:
                progress(number, len(indices), item)
            if result.cancelled:
                break
    finally:
        try:
            if session:
                session.__exit__(None, None, None)
        finally:
            result.remaining_indices = indices[len(result.items) - initial_count:]
            if phase:
                phase("保存清单", len(result.items) - initial_count, len(indices))
            save_manifest(result)
    return result


def continue_batch(previous, template, dataset, mapping, pattern, cancel=None, progress=None, *, phase=None):
    if not previous.remaining_indices:
        raise InputError("没有尚未开始的记录；失败部分请使用“重试失败项”。")
    return generate_batch(template, dataset, mapping, previous.folder.parent, pattern,
                          previous.remaining_indices, previous.export_pdf, cancel, progress, previous=previous, phase=phase)


def retry_batch(previous, template, dataset, mapping, pattern, cancel=None, progress=None, *, phase=None):
    result = BatchResult(previous.folder, [dict(item) for item in previous.items], False,
                         list(previous.indices), previous.export_pdf,
                         remaining_indices=list(previous.remaining_indices), used_stems=set(previous.used_stems))
    slots = [slot for slot, item in enumerate(result.items) if item["状态"] != "成功"]
    converter = pdf_converter() if result.export_pdf else None
    session = PdfSession(converter, cancel=cancel) if converter else None
    completed = 0
    try:
        for number, slot in enumerate(slots, 1):
            if cancel and cancel.is_set():
                result.cancelled = True
                break
            item = result.items[slot]
            try:
                values = values_for(template, mapping, dataset.records[result.indices[slot]])
                existing = result.folder / item["Word 文件"] if item["Word 文件"] else None
                if existing is None or not existing.is_file():
                    stem = file_stem(pattern, values, slot + 1)
                    destination = result.folder / f"{stem}.docx"
                    suffix = 2
                    while destination.exists():
                        destination = result.folder / f"{stem}_{suffix}.docx"
                        suffix += 1
                    if phase:
                        phase("填充 Word", number, len(slots))
                    template.render(values, destination)
                    existing = destination
                    item["Word 文件"] = existing.name
                    result.used_stems.add(existing.stem.casefold())
                if result.export_pdf:
                    if session is None:
                        raise InputError("未检测到可用 PDF 转换器。")
                    if phase:
                        phase("转换 PDF", number, len(slots))
                    item["PDF 文件"] = session.convert(existing).name
                item["状态"] = "成功"
                item["说明"] = "含空值字段" if any(not value.strip() for value in values.values()) else ""
            except Exception as exc:
                item["状态"] = "部分成功" if item["Word 文件"] else "失败"
                item["说明"] = str(exc)
                if isinstance(exc, ConversionCancelled):
                    result.cancelled = True
            completed = number
            if progress:
                progress(number, len(slots), item)
    finally:
        if session:
            session.__exit__(None, None, None)
        if phase:
            phase("保存清单", completed, len(slots))
        save_manifest(result)
    return result
